# Classification: INTERNAL
# AI-assisted (Cursor) — review before merge.
"""Task 6 — the six template download URLs. Requirements 4.1, 4.2, 1.2, 1.3."""

import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from importer import columns as columns_module
from importer.services import sheets

PAGES = sorted(columns_module.SHEETS)

pytestmark = pytest.mark.django_db


@pytest.fixture
def signed_in(client, django_user_model):
    user = django_user_model.objects.create_superuser(
        username="operator", password="synthetic-test-password", email="operator@example.com"
    )
    client.force_login(user)
    return client


@pytest.mark.parametrize("page", PAGES)
@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_both_formats_download_for_every_page(signed_in, page, fmt):
    response = signed_in.get(reverse(f"importer:{page}_template_{fmt}"))
    assert response.status_code == 200
    assert f'filename="netbox-{page}-template.{fmt}"' in response["Content-Disposition"]
    assert response.content


@pytest.mark.parametrize("page", PAGES)
def test_the_xlsx_download_is_a_real_workbook_with_the_right_sheet(signed_in, page):
    response = signed_in.get(reverse(f"importer:{page}_template_xlsx"))
    workbook = load_workbook(io.BytesIO(response.content))
    assert columns_module.SHEET_TITLES[page] in workbook.sheetnames


@pytest.mark.parametrize("page", PAGES)
def test_the_csv_download_headers_match_the_column_contract(signed_in, page):
    response = signed_in.get(reverse(f"importer:{page}_template_csv"))
    first_line = response.content.decode("utf-8-sig").splitlines()[0]
    expected = ",".join(spec.label for spec in columns_module.columns_for(page))
    assert first_line == expected


@pytest.mark.parametrize("page", PAGES)
def test_downloads_require_a_superuser(client, django_user_model, page):
    url = reverse(f"importer:{page}_template_csv")
    assert client.get(url).status_code == 302

    bystander = django_user_model.objects.create_user(
        username="bystander", password="synthetic-test-password"
    )
    client.force_login(bystander)
    assert client.get(url).status_code == 403


@pytest.mark.parametrize("page", PAGES)
def test_the_page_offers_both_downloads_and_documents_its_columns(signed_in, page):
    html = signed_in.get(reverse(f"importer:{page}")).content.decode()
    assert reverse(f"importer:{page}_template_csv") in html
    assert reverse(f"importer:{page}_template_xlsx") in html
    for spec in columns_module.columns_for(page):
        assert f"<code>{spec.label}</code>" in html


def test_hierarchy_page_explains_prepared_parents_and_opaque_rack_names(signed_in):
    html = signed_in.get(reverse("importer:hierarchy")).content.decode()
    assert "site and parent location must already exist" in html
    assert "Rack names are opaque join keys" in html
    assert "&lt;datacenter_slug&gt;" not in html


def test_a_downloaded_template_round_trips_through_the_view(signed_in):
    """The same guarantee as the unit test, but over HTTP: what the operator gets is what the
    parser accepts, apart from the deliberate refusal of the example row (Requirements 4.2, 4.5)."""
    response = signed_in.get(reverse("importer:hierarchy_template_csv"))
    payload = response.content
    with pytest.raises(sheets.SheetRejected) as raised:
        sheets.parse_upload(
            "hierarchy",
            filename="netbox-hierarchy-template.csv",
            stream=io.BytesIO(payload),
            size=len(payload),
        )
    assert [issue.code.value for issue in raised.value.issues] == [
        "EXAMPLE_ROW_NOT_REPLACED",
        "NO_DATA_ROWS",
    ]
